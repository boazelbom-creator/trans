"""Reading a community's aids.xlsx (PRD §7)."""

from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

SHEET_FORMAT = "פורמט"
SHEET_PHRASES = "ניסוחים נפוצים"
SHEET_TAGS = "תיוגים"
SHEET_KEYWORDS = "מילות מפתח לחיפוש"
SHEET_SOURCES = "ערוצי איתור משרות"

SOURCE_FACEBOOK = "facebook"
SOURCE_WHATSAPP = "whatsapp"
SOURCE_WEBSITE = "website"

SOURCE_TYPE_LABELS = {
    SOURCE_FACEBOOK: "פייסבוק",
    SOURCE_WHATSAPP: "ווטסאפ",
    SOURCE_WEBSITE: "אתר",
}


class ExcelError(Exception):
    """The Excel file is missing a sheet or column, or cannot be read. The message is shown to the user."""


@dataclass
class Phrase:
    description: str
    phrasing: str
    notes: str


@dataclass
class Tag:
    category: str
    name: str
    description: str
    notes: str


@dataclass
class Keyword:
    name: str
    details: str


@dataclass
class Source:
    name: str
    range: str
    url: str
    notes: str
    frequency: str

    @property
    def type(self) -> str:
        return source_type(self.url)


@dataclass
class AidsData:
    format_blocks: list[str] = field(default_factory=list)
    phrases: list[Phrase] = field(default_factory=list)
    tags: list[Tag] = field(default_factory=list)
    keywords: list[Keyword] = field(default_factory=list)
    sources: list[Source] = field(default_factory=list)


def source_type(url: str) -> str:
    host = (urlparse(url).hostname or "").lower()
    if host == "facebook.com" or host.endswith(".facebook.com") or host == "fb.com":
        return SOURCE_FACEBOOK
    if host == "whatsapp.com" or host.endswith(".whatsapp.com"):
        return SOURCE_WHATSAPP
    return SOURCE_WEBSITE


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


class _Sheet:
    """A sheet whose first row holds column headers."""

    def __init__(self, workbook, name: str):
        if name not in workbook.sheetnames:
            raise ExcelError(f"בקובץ האקסל חסר הגיליון \"{name}\".")
        self.name = name
        self.ws = workbook[name]
        first = next(self.ws.iter_rows(min_row=1, max_row=1), ())
        self.headers = {_text(c.value): c.column - 1 for c in first if _text(c.value)}

    def column(self, header: str, required: bool = True) -> int | None:
        if header in self.headers:
            return self.headers[header]
        if required:
            raise ExcelError(f"בגיליון \"{self.name}\" חסרה העמודה \"{header}\".")
        return None

    def rows(self):
        """Data rows (after the header row) as tuples of cells, skipping empty rows."""
        for row in self.ws.iter_rows(min_row=2):
            if any(_text(c.value) for c in row):
                yield row


def _cell(row, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return _text(row[index].value)


def _link(row, index: int) -> str:
    """The cell's hyperlink target, or its text if the text itself is a link."""
    if index >= len(row):
        return ""
    cell = row[index]
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target.strip()
    text = _text(cell.value)
    if text.lower().startswith(("http://", "https://", "www.")):
        return text if not text.lower().startswith("www.") else "https://" + text
    return ""


def _load_format(wb) -> list[str]:
    sheet = _Sheet(wb, SHEET_FORMAT)
    blocks = []
    for row in sheet.rows():
        for cell in row:
            text = _text(cell.value)
            if text:
                blocks.append(text)
    if not blocks:
        raise ExcelError(f"הגיליון \"{SHEET_FORMAT}\" ריק.")
    return blocks


def _load_phrases(wb) -> list[Phrase]:
    sheet = _Sheet(wb, SHEET_PHRASES)
    c_desc = sheet.column("תיאור דרישה [דוגמאות]")
    c_phr = sheet.column("ניסוח")
    c_notes = sheet.column("הערות", required=False)
    return [
        Phrase(_cell(r, c_desc), _cell(r, c_phr), _cell(r, c_notes))
        for r in sheet.rows()
        if _cell(r, c_desc) or _cell(r, c_phr)
    ]


def _load_tags(wb) -> list[Tag]:
    sheet = _Sheet(wb, SHEET_TAGS)
    c_tag = sheet.column("האשטג")
    c_desc = sheet.column("סוג התוכן", required=False)
    c_notes = sheet.column("הערות", required=False)
    # The category column has no header; it is filled only on the first row of each category.
    c_cat = sheet.column("קטגוריה", required=False)
    if c_cat is None and c_tag != 0:
        c_cat = 0
    tags = []
    category = ""
    for r in sheet.rows():
        if c_cat is not None and _cell(r, c_cat):
            category = _cell(r, c_cat)
        name = _cell(r, c_tag)
        if name:
            tags.append(Tag(category, name, _cell(r, c_desc), _cell(r, c_notes)))
    if not tags:
        raise ExcelError(f"בגיליון \"{SHEET_TAGS}\" אין תגיות.")
    return tags


def _load_keywords(wb) -> list[Keyword]:
    sheet = _Sheet(wb, SHEET_KEYWORDS)
    c_name = sheet.column("מילת חיפוש")
    c_details = sheet.column("פירוט", required=False)
    return [Keyword(_cell(r, c_name), _cell(r, c_details)) for r in sheet.rows() if _cell(r, c_name)]


def _load_sources(wb) -> list[Source]:
    sheet = _Sheet(wb, SHEET_SOURCES)
    c_name = sheet.column("שם ערוץ")
    c_link = sheet.column("לינק")
    c_range = sheet.column("טווח", required=False)
    c_notes = sheet.column("הערות", required=False)
    c_freq = sheet.column("תדירות", required=False)
    return [
        Source(_cell(r, c_name), _cell(r, c_range), _link(r, c_link), _cell(r, c_notes), _cell(r, c_freq))
        for r in sheet.rows()
        if _cell(r, c_name)
    ]


def load_aids(path: Path) -> AidsData:
    try:
        # Not read-only: read-only mode does not load hyperlinks.
        wb = openpyxl.load_workbook(path, data_only=True)
    except FileNotFoundError as e:
        raise ExcelError(f"קובץ האקסל חסר:\n{path}") from e
    except PermissionError as e:
        raise ExcelError(f"לא ניתן לפתוח את קובץ האקסל. אם הוא פתוח ב-Excel, סגרו אותו ונסו שוב.\n{path}") from e
    except Exception as e:  # openpyxl raises several types for damaged files
        raise ExcelError(f"לא ניתן לקרוא את קובץ האקסל:\n{path}\n{e}") from e
    try:
        return AidsData(
            format_blocks=_load_format(wb),
            phrases=_load_phrases(wb),
            tags=_load_tags(wb),
            keywords=_load_keywords(wb),
            sources=_load_sources(wb),
        )
    finally:
        wb.close()
